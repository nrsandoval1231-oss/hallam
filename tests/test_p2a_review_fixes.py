"""Fixes from the Pass 2A review: identity truth levels, run safety, uploads, demotion."""

from __future__ import annotations

import io
import shutil
import sqlite3
import subprocess
import sys
import threading
import time
from datetime import UTC, datetime, timedelta
from pathlib import Path
from types import SimpleNamespace

import pytest
from docx import Document
from fastapi.testclient import TestClient
from sabra_api.main import app
from sabra_domain import models
from sabra_domain.db import Database
from sabra_domain.errors import ValidationFailure
from sabra_domain.state import can_transition_job
from sabra_engine import runs as runs_module
from sabra_engine.answer_memory import AnswerMemoryService
from sabra_engine.candidate_intake import extract_resume_text, save_owner_identity
from sabra_engine.cli import app as cli_app
from sabra_engine.context import EngineContext
from sabra_engine.intelligence import IntelligenceService, pass_opportunity
from sabra_engine.research import ResearchService, refusal_for_url
from sabra_engine.runs import RunService, pid_alive
from sqlalchemy import select
from test_p2a_support import mock_http, seed_candidate, use_registry
from typer.testing import CliRunner


@pytest.fixture()
def world(tmp_path, monkeypatch, fixture_resume_text):
    use_registry(tmp_path, monkeypatch)
    return seed_candidate(fixture_resume_text)


def _service(ctx, **kwargs) -> RunService:
    return RunService(ctx, discovery_http=mock_http(), research_http=mock_http(), **kwargs)


def _candidate_json() -> dict:
    return TestClient(app).get("/api/candidate").json()


# ---------------------------------------------------------------------------
# B1: the identity form only records what the owner changed or confirmed
# ---------------------------------------------------------------------------


def _with_resume(fixture_resume_text: str) -> TestClient:
    client = TestClient(app)
    response = client.post(
        "/api/candidate/resume",
        files={"resume": ("resume.txt", fixture_resume_text.encode(), "text/plain")},
    )
    assert response.status_code == 200
    return client


def test_saving_one_field_leaves_every_other_value_at_its_previous_level(
    fixture_resume_text,
) -> None:
    client = _with_resume(fixture_resume_text)
    ctx = EngineContext.create()
    try:
        candidate_id = client.get("/api/candidate").json()["candidate_id"]
        # An answer recorded but never verified.
        AnswerMemoryService(ctx).record("street_address", "9 Unverified Rd")
        with ctx.read_repos() as repos:
            version_before = repos.candidates.latest_profile_row(candidate_id).version
    finally:
        ctx.close()
    before = client.get("/api/candidate").json()
    levels_before = before["identity_levels"]
    assert levels_before["email"] == "INFERRED" and levels_before["phone"] == "INFERRED"
    assert levels_before["street_address"] == "INFERRED"
    assert levels_before["headline"] == "INFERRED"

    saved = client.put("/api/candidate/identity", json={"headline": "Office professional"})
    assert saved.status_code == 200
    assert saved.json()["version"] == version_before + 1
    after = client.get("/api/candidate").json()
    assert after["identity_levels"]["headline"] == "VERIFIED"
    for field in ("email", "phone", "city", "state", "linkedin_url", "street_address"):
        assert after["identity_levels"][field] == levels_before[field], field
        assert after["identity"][field] == before["identity"][field], field
    assert after["provenance"]["identity.email"]["level"] == "INFERRED"
    assert after["provenance"]["employment"]["level"] == "INFERRED"
    assert after["profile"]["answer_memory"]["street_address"]["level"] == "INFERRED"


def test_submitting_a_stored_value_confirms_it_and_records_where_it_came_from(
    fixture_resume_text,
) -> None:
    client = _with_resume(fixture_resume_text)
    before = client.get("/api/candidate").json()
    email = before["identity"]["email"]
    phone_level = before["identity_levels"]["phone"]
    assert client.put("/api/candidate/identity", json={"email": email}).status_code == 200
    after = client.get("/api/candidate").json()
    assert after["identity_levels"]["email"] == "VERIFIED"
    record = after["provenance"]["identity.email"]
    assert record["evidence"][0] == "owner_identity"
    assert "confirmed_from:master_resume" in record["evidence"]
    assert "INFERRED" in record["note"]
    assert after["identity"]["email"] == email
    assert after["identity_levels"]["phone"] == phone_level == "INFERRED"
    # Confirming a value that is already VERIFIED changes nothing and writes no new version.
    version = client.put("/api/candidate/identity", json={"email": email}).json()["version"]
    again = client.put("/api/candidate/identity", json={"email": email}).json()["version"]
    assert again == version


def test_confirming_an_unverified_street_address_makes_it_verified_with_its_origin() -> None:
    client = TestClient(app)
    client.put("/api/candidate/identity", json={"first_name": "Maria"})
    ctx = EngineContext.create()
    try:
        AnswerMemoryService(ctx).record("street_address", "9 Unverified Rd")
    finally:
        ctx.close()
    assert client.get("/api/candidate").json()["identity_levels"]["street_address"] == "INFERRED"
    client.put("/api/candidate/identity", json={"street_address": "9 Unverified Rd"})
    answer = client.get("/api/candidate").json()["profile"]["answer_memory"]["street_address"]
    assert answer["level"] == "VERIFIED"
    assert "confirmed_from:owner_recorded" in answer["evidence"]
    assert "INFERRED" in answer["note"]


def test_the_form_sends_only_changed_or_confirmed_values() -> None:
    from test_p2a_html_safety import render

    payload = {
        "identity": {"email": "a@b.test", "phone": "555", "city": "Plano"},
        "identity_levels": {"email": "INFERRED", "phone": "VERIFIED", "city": "INFERRED"},
        "sections": [],
        "resumes": [],
    }
    from sabra_api.main import candidate_page

    expression = """
const f = new FormData({email:'a@b.test', phone:'555', city:'Plano', state:'TX',
  headline:'  New headline ', last_name:'   ', confirm_city:'on'});
document.querySelector('#out').textContent = JSON.stringify(collectIdentity(f, prefilled));
"""
    nodes = render(candidate_page(), payload, expression)
    import json

    sent = json.loads(nodes["#out"]["textContent"])
    # Unchanged prefilled email/phone are not sent; a changed, a new and a confirmed value are.
    assert sent == {"city": "Plano", "state": "TX", "headline": "New headline"}
    # The level shows beside each box, and a verified (or empty) value has no confirm box.
    assert "INFERRED" in nodes["#lvl-email"]["innerHTML"]
    assert "VERIFIED" in nodes["#lvl-phone"]["innerHTML"]
    assert nodes["#confirm-email"]["hidden"] is False
    assert nodes["#confirm-phone"]["hidden"] is True
    assert nodes["#confirm-state"]["hidden"] is True


def test_saving_with_nothing_changed_sends_nothing() -> None:
    from sabra_api.main import candidate_page
    from test_p2a_html_safety import render

    payload = {
        "identity": {"email": "a@b.test"},
        "identity_levels": {"email": "INFERRED"},
        "sections": [],
        "resumes": [],
    }
    expression = (
        "document.querySelector('#identity-form').listeners.submit("
        "{preventDefault(){}, target:{email:'a@b.test'}})"
    )
    nodes = render(candidate_page(), payload, expression)
    assert "Nothing to save" in nodes["#identity-line"]["textContent"]


def test_every_identity_input_has_a_level_slot_and_a_confirm_box() -> None:
    html = TestClient(app).get("/candidate").text
    for name in ("first_name", "email", "street_address", "headline"):
        assert f'id="lvl-{name}"' in html and f'name="confirm_{name}"' in html


# ---------------------------------------------------------------------------
# S1: a role passed between the assess batches and the shortlist step
# ---------------------------------------------------------------------------


def test_a_pass_during_the_pass_does_not_fail_or_shortlist_the_role(world) -> None:
    candidate_id, profile_id = world
    ctx = EngineContext.create()
    try:
        _service(ctx).run_sync("discover", False)
        passed: list[str] = []

        def progress(payload: dict) -> None:
            if payload["stage"] == "rank" and not passed:
                with ctx.read_repos() as repos:
                    qualified = [
                        o.id for o in repos.opportunities.list_for_candidate(candidate_id)
                        if o.status == "QUALIFIED"
                    ]
                passed.append(qualified[0])
                pass_opportunity(ctx, qualified[0])

        result = IntelligenceService(ctx).run(
            candidate_id=candidate_id, search_profile_id=profile_id, progress=progress
        )
        assert passed, "the callback never saw the rank stage"
        assert result.shortlisted == 0
        with ctx.read_repos() as repos:
            assert repos.opportunities.shortlist(candidate_id) == []
            row = repos.opportunities.get(passed[0])
            assert row.status == "REJECTED" and row.rejection["reasons"] == ["OWNER_PASSED"]
        assert all(entry["opportunity_id"] != passed[0] for entry in result.top)
    finally:
        ctx.close()


# ---------------------------------------------------------------------------
# S2: a reaped run that is still alive must not overwrite the verdict
# ---------------------------------------------------------------------------


def test_a_run_ended_by_someone_else_stops_and_keeps_its_failed_record(world) -> None:
    ctx = EngineContext.create()
    flipped: list[bool] = []

    def reaper(_progress: dict) -> None:
        if not flipped:
            flipped.append(True)
            with ctx.repos() as repos:
                row = repos.pipeline_runs.list_recent(1)[0]
                row.status = "FAILED"
                row.error = "abandoned: reaped by the test"
                row.finished_at = datetime.now(UTC)

    try:
        record = _service(ctx, on_update=reaper).run_sync("full", False)
    finally:
        ctx.close()
    assert record["status"] == "FAILED"
    assert record["error"] == "abandoned: reaped by the test"
    assert "rank" not in record["progress"]["stages"], "the run must stop, not carry on"
    ctx = EngineContext.create()
    try:
        with ctx.read_repos() as repos:
            events = [e.event_type for e in repos.events.recent(200)]
        assert "PIPELINE_RUN_COMPLETED" not in events
    finally:
        ctx.close()


def test_the_reaper_asks_whether_the_owner_process_is_alive() -> None:
    now = datetime.now(UTC)

    def row(owner: dict | None, age: timedelta):
        return SimpleNamespace(
            progress={"owner": owner} if owner else {}, updated_at=now - age, created_at=now - age
        )

    import socket

    host = socket.gethostname()
    finished = subprocess.Popen([sys.executable, "-c", "pass"])
    finished.wait()
    assert pid_alive(runs_module.os.getpid()) is True
    assert pid_alive(finished.pid) is False

    dead = {"pid": finished.pid, "host": host}
    live = {"pid": runs_module.os.getpid(), "host": host}
    # Dead owner on this machine: reaped at once, however fresh.
    assert runs_module._is_dead(row(dead, timedelta(seconds=5))) is True
    # Live owner: a quiet step is not death, even after the old 15-minute mark...
    assert runs_module._is_dead(row(live, timedelta(minutes=30))) is False
    # ...but a run silent for hours is wedged (or the pid was reused).
    assert runs_module._is_dead(row(live, runs_module.HARD_STALE_AFTER * 2)) is True
    # Another machine, or no owner: only the clock can say.
    other = {"pid": 1, "host": "some-other-machine"}
    assert runs_module._is_dead(row(other, timedelta(minutes=30))) is False
    assert runs_module._is_dead(row(other, runs_module.STALE_AFTER * 2)) is True
    assert runs_module._is_dead(row(None, runs_module.STALE_AFTER * 2)) is True
    assert runs_module._is_dead(row(None, timedelta(minutes=30))) is False


def test_a_run_whose_process_is_gone_is_reaped_so_the_next_run_can_start(world) -> None:
    finished = subprocess.Popen([sys.executable, "-c", "pass"])
    finished.wait()
    ctx = EngineContext.create()
    try:
        dead = RunService(ctx).create("full")
        with ctx.repos() as repos:
            row = repos.pipeline_runs.get(dead)
            row.status = "RUNNING"
            import socket

            row.progress = {"owner": {"pid": finished.pid, "host": socket.gethostname()}}
        replacement = RunService(ctx).create("rank")
        assert replacement != dead
        assert RunService(ctx).get(dead)["status"] == "FAILED"
        # And a live owner is not reaped by a fresh request.
        with pytest.raises(ValidationFailure, match="already in progress"):
            RunService(ctx).create("rank")
    finally:
        ctx.close()


# ---------------------------------------------------------------------------
# S3: one active run for the whole database, and serialised write scopes
# ---------------------------------------------------------------------------


def test_the_database_refuses_a_second_active_run(world) -> None:
    ctx = EngineContext.create()
    try:
        with ctx.repos() as repos:
            repos.pipeline_runs.create(kind="full", prepare=False, candidate_id="c", search_profile_id=None)
        with pytest.raises(Exception) as raised, ctx.repos() as repos:
            repos.pipeline_runs.create(
                kind="rank", prepare=False, candidate_id="c", search_profile_id=None
            )
        assert "UNIQUE" in str(raised.value).upper()
        # A finished run does not count, so history can grow.
        with ctx.repos() as repos:
            repos.pipeline_runs.unfinished()[0].status = "COMPLETED"
        with ctx.repos() as repos:
            repos.pipeline_runs.create(kind="rank", prepare=False, candidate_id="c", search_profile_id=None)
    finally:
        ctx.close()


def test_create_maps_the_database_refusal_to_already_in_progress(world, monkeypatch) -> None:
    ctx = EngineContext.create()
    try:
        RunService(ctx).create("full")
        # Another process got there between our check and our insert: hide the row from the
        # check so only the index can object.
        monkeypatch.setattr(
            "sabra_domain.repositories.PipelineRunRepository.unfinished", lambda self: []
        )
        with pytest.raises(ValidationFailure, match="already in progress"):
            RunService(ctx).create("rank")
    finally:
        ctx.close()


def test_two_threads_starting_a_run_at_once_start_exactly_one(world) -> None:
    outcomes: list[str] = []
    barrier = threading.Barrier(2)

    def start() -> None:
        ctx = EngineContext.create()
        try:
            barrier.wait(10)
            try:
                RunService(ctx).create("full")
                outcomes.append("started")
            except ValidationFailure:
                outcomes.append("refused")
        finally:
            ctx.close()

    threads = [threading.Thread(target=start) for _ in range(2)]
    for t in threads:
        t.start()
    for t in threads:
        t.join(30)
    assert sorted(outcomes) == ["refused", "started"]


def test_an_existing_runs_table_gets_the_index_when_the_schema_is_created(settings) -> None:
    settings.ensure_dirs()
    database = Database(settings.database_url)
    try:
        database.create_all()
        with database.engine.begin() as connection:
            connection.exec_driver_sql("DROP INDEX uq_runs_one_active")
        database.create_all()
        with database.engine.connect() as connection:
            names = {
                row[0]
                for row in connection.exec_driver_sql(
                    "SELECT name FROM sqlite_master WHERE type = 'index'"
                )
            }
        assert "uq_runs_one_active" in names
    finally:
        database.dispose()


def test_write_scopes_serialise_so_read_then_write_never_loses_an_update(world) -> None:
    errors: list[BaseException] = []

    def bump() -> None:
        ctx = EngineContext.create()
        try:
            for _ in range(15):
                with ctx.repos() as repos:
                    value = (repos.app_state.get("counter") or {"v": 0})["v"]
                    time.sleep(0.002)  # a window in which the other thread would interleave
                    repos.app_state.set("counter", {"v": value + 1})
        except BaseException as exc:  # noqa: BLE001
            errors.append(exc)
        finally:
            ctx.close()

    threads = [threading.Thread(target=bump) for _ in range(2)]
    for t in threads:
        t.start()
    for t in threads:
        t.join(60)
    assert not errors, errors
    ctx = EngineContext.create()
    try:
        with ctx.read_repos() as repos:
            assert repos.app_state.get("counter") == {"v": 30}
    finally:
        ctx.close()


def test_a_write_scope_begins_immediate_and_a_read_scope_does_not(world, tmp_path) -> None:
    ctx = EngineContext.create()
    db_file = ctx.settings.database_url.split("///", 1)[1]
    try:
        with ctx.repos():
            # Another connection cannot write while this scope is open, even before it writes.
            other = sqlite3.connect(db_file, timeout=0.2)
            with pytest.raises(sqlite3.OperationalError):
                other.execute("BEGIN IMMEDIATE")
            other.close()
        with ctx.read_repos():
            other = sqlite3.connect(db_file, timeout=0.2)
            other.execute("BEGIN IMMEDIATE")
            other.execute("ROLLBACK")
            other.close()
        with ctx.db.session(immediate=False):
            other = sqlite3.connect(db_file, timeout=0.2)
            other.execute("BEGIN IMMEDIATE")
            other.execute("ROLLBACK")
            other.close()
    finally:
        ctx.close()


def test_plain_discover_and_rank_refuse_while_a_run_is_active(world) -> None:
    ctx = EngineContext.create()
    try:
        RunService(ctx).create("full")
    finally:
        ctx.close()
    runner = CliRunner()
    for command in ("discover", "rank"):
        result = runner.invoke(cli_app, [command])
        assert result.exit_code != 0, command
        assert isinstance(result.exception, ValidationFailure), command
        assert "already in progress" in str(result.exception)


# ---------------------------------------------------------------------------
# S4: interruption
# ---------------------------------------------------------------------------


def test_ctrl_c_ends_the_run_failed_and_still_propagates(world, monkeypatch) -> None:
    def interrupted(self, **kwargs):
        raise KeyboardInterrupt

    monkeypatch.setattr("sabra_engine.discovery.DiscoveryService.run", interrupted)
    ctx = EngineContext.create()
    try:
        service = _service(ctx)
        run_id = service.create("full")
        with pytest.raises(KeyboardInterrupt):
            service.execute(run_id)
        record = RunService(ctx).get(run_id)
        assert record["status"] == "FAILED"
        assert "interrupted (KeyboardInterrupt)" in record["error"]
        assert record["finished_at"]
        # The row no longer blocks the next run.
        assert RunService(ctx).create("rank")
    finally:
        ctx.close()


# ---------------------------------------------------------------------------
# S5: "new high-fit" follows the shortlist event
# ---------------------------------------------------------------------------


def test_a_rank_run_after_a_discover_run_reports_its_new_roles(world) -> None:
    ctx = EngineContext.create()
    try:
        _service(ctx).run_sync("discover", False)
        client = TestClient(app)
        assert client.get("/api/today").json()["counts"]["new_high_fit"] == 0
        time.sleep(0.01)
        _service(ctx).run_sync("rank", False)
        assert client.get("/api/today").json()["counts"]["new_high_fit"] == 1
        time.sleep(0.01)
        _service(ctx).run_sync("rank", False)  # nothing new this time
        assert client.get("/api/today").json()["counts"]["new_high_fit"] == 0
    finally:
        ctx.close()


# ---------------------------------------------------------------------------
# S6: demotion goes through the graph
# ---------------------------------------------------------------------------


class _EmptyShortlist:
    """A ranker that ranks as usual but selects nobody, so everyone is demoted."""

    def __init__(self) -> None:
        from sabra_scoring.ranking import Ranker

        self._real = Ranker()

    def rank(self, assessments):
        return self._real.rank(assessments)

    def select_shortlist(self, ranked, *, limit, eligible=None):
        return []


def test_demotion_uses_declared_edges_moves_the_postings_and_leaves_an_event(world) -> None:
    candidate_id, profile_id = world
    assert can_transition_job(models_status("SHORTLISTED"), models_status("QUALIFIED"))
    assert can_transition_job(models_status("RESEARCHED"), models_status("QUALIFIED"))
    ctx = EngineContext.create()
    try:
        _service(ctx).run_sync("full", False)  # leaves the role RESEARCHED
        with ctx.read_repos() as repos:
            opportunity = repos.opportunities.shortlist(candidate_id)[0]
            assert opportunity.status == "RESEARCHED"
            opportunity_id = opportunity.id
        IntelligenceService(ctx, ranker=_EmptyShortlist()).run(
            candidate_id=candidate_id, search_profile_id=profile_id
        )
        with ctx.read_repos() as repos:
            row = repos.opportunities.get(opportunity_id)
            assert row.status == "QUALIFIED" and row.shortlisted is False
            assert "demoted" in row.shortlist_reason
            jobs = repos.opportunities.list_jobs(opportunity_id)
            assert jobs and all(job.status == "QUALIFIED" for job in jobs)
            events = repos.events.for_entity("opportunity", opportunity_id)
            demotions = [
                e for e in events
                if e.event_type == "JOB_SCORED" and e.payload.get("stage") == "demoted_from_shortlist"
            ]
            assert demotions and demotions[0].payload["from"] == "RESEARCHED"
    finally:
        ctx.close()


def models_status(name: str):
    from sabra_domain.enums import JobStatus

    return JobStatus(name)


# ---------------------------------------------------------------------------
# S7: uploads
# ---------------------------------------------------------------------------


def test_utf16_text_is_decoded_and_binary_or_unknown_types_are_refused() -> None:
    text_value = "Jane Doe\nOffice Manager"
    assert extract_resume_text("a.txt", text_value.encode("utf-16")) == (text_value, "text")
    assert extract_resume_text("a.txt", text_value.encode("utf-16-be").join([b"\xfe\xff", b""]))[0]
    assert extract_resume_text("notes.MD", b"# Jane\n\nOffice") == ("# Jane\n\nOffice", "text")
    with pytest.raises(ValidationFailure, match="binary"):
        extract_resume_text("a.txt", text_value.encode("utf-16-le"))  # no BOM: NULs remain
    with pytest.raises(ValidationFailure, match="binary"):
        extract_resume_text("a.txt", b"abc\x00def")
    for name in ("resume.rtf", "resume.odt", "resume", "resume.exe", None):
        with pytest.raises(ValidationFailure, match="only .pdf"):
            extract_resume_text(name, b"{\\rtf1 hello}")


def test_the_upload_route_answers_422_for_an_rtf_and_a_utf16_txt_is_stored_readably() -> None:
    client = TestClient(app)
    rejected = client.post(
        "/api/candidate/resume", files={"resume": ("resume.rtf", b"{\\rtf1 hi}", "text/rtf")}
    )
    assert rejected.status_code == 422
    stored = client.post(
        "/api/candidate/resume",
        files={"resume": ("resume.txt", "Jane Doe\nOffice Manager".encode("utf-16"), "text/plain")},
    )
    assert stored.status_code == 200, stored.text
    ctx = EngineContext.create()
    try:
        with ctx.read_repos() as repos:
            resume = repos.session.execute(select(models.Resume)).scalars().one()
            assert resume.content_text == "Jane Doe\nOffice Manager"
            assert "\x00" not in resume.content_text
    finally:
        ctx.close()


def test_docx_headers_and_footers_are_read_once_around_the_body() -> None:
    document = Document()
    first = document.sections[0]
    first.header.paragraphs[0].text = "Jane Doe | jane@example.test | (214) 555-0100"
    first.footer.paragraphs[0].text = "References on request"
    document.add_paragraph("EXPERIENCE")
    document.add_section()  # linked headers and footers: the same text, not repeated
    document.add_paragraph("EDUCATION")
    buffer = io.BytesIO()
    document.save(buffer)
    text_value, source_format = extract_resume_text("r.docx", buffer.getvalue())
    assert source_format == "docx"
    assert text_value.split("\n") == [
        "Jane Doe | jane@example.test | (214) 555-0100",
        "EXPERIENCE",
        "EDUCATION",
        "References on request",
    ]


# ---------------------------------------------------------------------------
# S8: the scheduled runner
# ---------------------------------------------------------------------------

SCRIPTS = Path(__file__).parents[1] / "scripts"


def test_the_runner_logs_utf8_and_cleans_its_inputs() -> None:
    text_value = (SCRIPTS / "run-sabra-daily.ps1").read_text(encoding="utf-8")
    assert "[Console]::OutputEncoding = [System.Text.Encoding]::UTF8" in text_value
    assert "2>&1 | ForEach-Object" in text_value and "Out-File" in text_value
    assert text_value.count("-Encoding utf8") >= 3
    assert "*>>" not in text_value
    assert "TrimEnd" in text_value and "throw" in text_value
    assert 'TrimEnd("\\", "/")' in (SCRIPTS / "schedule-windows.ps1").read_text(encoding="utf-8")


@pytest.mark.skipif(
    shutil.which("powershell") is None and shutil.which("pwsh") is None,
    reason="PowerShell is not installed",
)
def test_the_runner_stops_before_running_anything_when_the_env_script_is_missing(tmp_path) -> None:
    shell = shutil.which("powershell") or shutil.which("pwsh")
    result = subprocess.run(
        [
            shell, "-NoProfile", "-ExecutionPolicy", "Bypass", "-File",
            str(SCRIPTS / "run-sabra-daily.ps1"), "-PrivateDir", str(tmp_path) + "\\",
        ],
        capture_output=True, text=True, timeout=60,
    )
    assert result.returncode != 0
    assert "hallam-env.ps1" in result.stdout + result.stderr
    assert not (tmp_path / "var" / "logs").exists(), "nothing may have been run or logged"


# ---------------------------------------------------------------------------
# N4, N6, N7, N8
# ---------------------------------------------------------------------------


def test_passing_an_expired_role_says_so_and_a_filtered_role_keeps_its_reasons(world) -> None:
    candidate_id, profile_id = world
    ctx = EngineContext.create()
    try:
        _service(ctx).run_sync("discover", False)
        IntelligenceService(ctx).run(candidate_id=candidate_id, search_profile_id=profile_id)
        with ctx.read_repos() as repos:
            rows = repos.opportunities.list_for_candidate(candidate_id)
            rejected = next(o for o in rows if o.status == "REJECTED")
            qualified = next(o for o in rows if o.status != "REJECTED")
            earlier = list(rejected.rejection["reasons"])
            earlier_detail = dict(rejected.rejection["detail"])
        pass_opportunity(ctx, rejected.id)
        with ctx.read_repos() as repos:
            now = repos.opportunities.get(rejected.id).rejection
        assert now["reasons"] == [*earlier, "OWNER_PASSED"]
        assert earlier_detail.items() <= now["detail"].items()

        with ctx.repos() as repos:
            repos.opportunities.set_status(qualified.id, models_status("EXPIRED"))
        client = TestClient(app)
        response = client.post(f"/api/opportunities/{qualified.id}/pass")
        assert response.status_code == 409
        assert "expired" in response.json()["detail"]
    finally:
        ctx.close()


def test_prepare_target_counts_only_roles_that_have_no_package(prepared) -> None:
    from conftest_app import build_job
    from sabra_domain.enums import JobStatus

    with prepared.ctx.repos() as repos:
        profile = repos.search_profiles.get(prepared.search_profile_id)
        profile.prepare_target = 1
        repos.search_profiles.upsert(profile)
        row, _, _ = repos.jobs.upsert_canonical(
            build_job(source_job_id="9002", title="Territory Manager - Dallas"), raw_payload={}
        )
        row.status = str(JobStatus.NORMALIZED)
        fresh = repos.opportunities.create(
            candidate_id=prepared.candidate_id,
            search_profile_id=prepared.search_profile_id,
            canonical_key=row.canonical_key,
            title=row.title,
            company=row.company,
            location=row.location,
            primary_job_id=row.id,
            source_ids=[{"source": row.source, "source_job_id": row.source_job_id, "url": row.source_url}],
        )
        row.opportunity_id = fresh.id
        fresh.status = str(JobStatus.SHORTLISTED)
        fresh.shortlisted = True
        fresh.rank_score = 0.5  # ranks below the role that already has a package
        fresh_id = fresh.id
    service = RunService(prepared.ctx)
    service._progress = {"stages": {}, "warnings": [], "sources": []}
    service._stage_prepare(prepared.candidate_id, prepared.search_profile_id)
    stage = service._progress["stages"]["prepare"]
    assert (stage["prepared"], stage["failed"], stage["skipped_existing"]) == (1, 0, 1)
    with prepared.ctx.read_repos() as repos:
        assert repos.applications.for_opportunity(fresh_id) is not None
        assert repos.packages.latest_for_opportunity(prepared.opportunity_id).revision == 1


def test_full_name_follows_an_actual_change_to_either_name_part(fixture_resume_text) -> None:
    client = _with_resume(fixture_resume_text)
    original = client.get("/api/candidate").json()["profile"]["identity"]["full_name"]
    client.put("/api/candidate/identity", json={"last_name": "Rivera"})
    assert client.get("/api/candidate").json()["profile"]["identity"]["full_name"] == original
    client.put("/api/candidate/identity", json={"first_name": "Jordan"})
    assert client.get("/api/candidate").json()["profile"]["identity"]["full_name"] == "Jordan Rivera"
    client.put("/api/candidate/identity", json={"last_name": "Rivera-Smith"})
    assert (
        client.get("/api/candidate").json()["profile"]["identity"]["full_name"]
        == "Jordan Rivera-Smith"
    )
    # Re-submitting an unchanged part does not rebuild the name.
    ctx = EngineContext.create()
    try:
        candidate_id = client.get("/api/candidate").json()["candidate_id"]
        save_owner_identity(ctx, candidate_id, {"first_name": "Jordan"})
        with ctx.read_repos() as repos:
            assert repos.candidates.active_profile(candidate_id).identity.full_name == (
                "Jordan Rivera-Smith"
            )
    finally:
        ctx.close()


@pytest.mark.parametrize(
    "url",
    [
        "file:///etc/passwd",
        "ftp://example.com/",
        "http://localhost/",
        "http://127.0.0.1/",
        "http://[::1]/",
        "http://10.0.0.5/",
        "http://192.168.1.1/",
        "http://172.16.0.1/",
        "http://169.254.169.254/latest/meta-data",
        "http://0.0.0.0/",
        "http://printer.local/",
        "http://intranet/",
        "http://db.internal/",
        "https:///nohost",
    ],
)
def test_non_public_or_non_http_urls_are_refused(url) -> None:
    assert refusal_for_url(url) is not None


def test_public_urls_pass_and_a_public_name_pointing_inside_is_refused(monkeypatch) -> None:
    assert refusal_for_url("https://revance.com") is None
    assert refusal_for_url("http://93.184.216.34/") is None
    monkeypatch.setattr(
        "socket.getaddrinfo", lambda host, port: [(2, 1, 6, "", ("10.1.2.3", 0))]
    )
    assert "non-public" in refusal_for_url("https://rebind.example.com", resolve=True)
    monkeypatch.setattr(
        "socket.getaddrinfo", lambda host, port: [(2, 1, 6, "", ("93.184.216.34", 0))]
    )
    assert refusal_for_url("https://fine.example.com", resolve=True) is None


def test_research_never_requests_a_non_public_company_site(world) -> None:
    candidate_id, profile_id = world
    ctx = EngineContext.create()
    try:
        _service(ctx).run_sync("discover", False)
        IntelligenceService(ctx).run(candidate_id=candidate_id, search_profile_id=profile_id)
        with ctx.repos() as repos:
            for job in repos.session.execute(select(models.Job)).scalars():
                job.company_domain = "169.254.169.254"
        requests: list[str] = []
        summary = ResearchService(ctx, http=mock_http(requests=requests)).research_fresh(
            candidate_id
        )
        assert summary["researched"] == 1 and summary["failed"] == 0
        assert requests == [], "a metadata-service address must never be requested"
        with ctx.read_repos() as repos:
            opportunity = repos.opportunities.shortlist(candidate_id)[0]
            packet = repos.research.for_opportunity(opportunity.id).payload
        assert "no company description could be captured" in " ".join(packet["inferences"])
    finally:
        ctx.close()


def test_the_redirect_hook_of_the_real_client_refuses_private_targets(world, monkeypatch) -> None:
    # The service builds its own client when none is injected; its request hook runs on every
    # request, so a redirect to a private address is refused before anything is sent.
    monkeypatch.setattr(
        "socket.getaddrinfo", lambda host, port: [(2, 1, 6, "", ("10.9.9.9", 0))]
    )
    ctx = EngineContext.create()
    try:
        service = ResearchService(ctx)
        description, sources, note = service._fetch_site("https://looks-public.example.com")
        assert description is None and sources == []
        assert "non-public" in (note or "")
    finally:
        ctx.close()


# ---------------------------------------------------------------------------
# N10 and N11
# ---------------------------------------------------------------------------


def test_the_searches_form_says_a_blank_age_limit_means_no_limit() -> None:
    assert "Leave blank for no limit" in TestClient(app).get("/searches").text


def test_owner_identity_wins_when_the_resume_comes_first(fixture_resume_text) -> None:
    client = _with_resume(fixture_resume_text)
    before = client.get("/api/candidate").json()
    assert before["identity"]["email"] == "jordan.rivera@example.test"
    assert before["provenance"]["identity.email"]["level"] == "INFERRED"
    client.put("/api/candidate/identity", json={"email": "later@example.test"})
    after = client.get("/api/candidate").json()
    assert after["identity"]["email"] == "later@example.test"
    assert after["provenance"]["identity.email"]["level"] == "VERIFIED"
    # A later resume upload cannot take it back.
    client.post(
        "/api/candidate/resume",
        files={"resume": ("again.txt", fixture_resume_text.encode(), "text/plain")},
    )
    again = client.get("/api/candidate").json()
    assert again["identity"]["email"] == "later@example.test"
