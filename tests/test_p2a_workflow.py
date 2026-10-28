"""Today, the Candidate page API, Settings and Pass."""

from __future__ import annotations

import io
from datetime import UTC, datetime, timedelta

import pytest
from docx import Document
from fastapi.testclient import TestClient
from sabra_api.main import app
from sabra_domain.enums import ApplicationStatus, EventType, PackageStatus
from sabra_domain.errors import ValidationFailure
from sabra_engine.candidate_intake import extract_resume_text, save_owner_identity
from sabra_engine.context import EngineContext
from sabra_engine.documents import DocumentService
from sabra_engine.intelligence import IntelligenceService, pass_opportunity
from sabra_engine.runs import RunService
from test_p2a_support import mock_http, seed_candidate, use_registry


@pytest.fixture()
def client_for_prepared(prepared, monkeypatch) -> TestClient:
    monkeypatch.setattr("sabra_engine.context.get_settings", lambda: prepared.ctx.settings)
    return TestClient(app)


# ---------------------------------------------------------------------------
# Today
# ---------------------------------------------------------------------------


def test_today_on_an_empty_database_is_all_zero() -> None:
    body = TestClient(app).get("/api/today").json()
    assert body["counts"] == {
        "ready_for_approval": 0,
        "needs_information": 0,
        "handoffs_waiting": 0,
        "new_high_fit": 0,
        "employer_responses": 0,
    }
    assert body["has_resume"] is False
    assert body["last_run"] is None and body["active_run"] is None
    assert body["automation_dry_run"] is True


def _set_package(prepared, **fields) -> None:
    with prepared.ctx.repos() as repos:
        package = repos.packages.get(prepared.package_id)
        for key, value in fields.items():
            setattr(package, key, value)


def test_today_counts_packages_ready_and_packages_needing_information(
    prepared, client_for_prepared
) -> None:
    _set_package(prepared, status=PackageStatus.COMPLETE.value, truth_clean=True)
    counts = client_for_prepared.get("/api/today").json()["counts"]
    assert counts["ready_for_approval"] == 1 and counts["needs_information"] == 0

    _set_package(prepared, status=PackageStatus.INCOMPLETE.value, unresolved_count=3)
    body = client_for_prepared.get("/api/today").json()
    assert body["counts"]["ready_for_approval"] == 0
    assert body["counts"]["needs_information"] == 1
    assert body["open_questions"] == 3
    assert body["has_resume"] is True


def test_today_counts_a_handoff_waiting_for_a_report(prepared, client_for_prepared) -> None:
    with prepared.ctx.repos() as repos:
        repos.applications.set_status(prepared.application_id, ApplicationStatus.APPROVED)
        repos.applications.set_status(prepared.application_id, ApplicationStatus.APPLYING)
        repos.events.append("application", prepared.application_id, EventType.APPLICATION_STARTED, {})
        repos.events.append(
            "application",
            prepared.application_id,
            EventType.HANDED_TO_HUMAN,
            {"state": "needs_report"},
        )
    counts = client_for_prepared.get("/api/today").json()["counts"]
    assert counts["handoffs_waiting"] == 1
    # An APPLYING row with nothing waiting on the person is not counted.
    with prepared.ctx.repos() as repos:
        repos.events.append("application", prepared.application_id, EventType.APPLICATION_STARTED, {})
    assert client_for_prepared.get("/api/today").json()["counts"]["handoffs_waiting"] == 0


def test_new_high_fit_counts_roles_shortlisted_since_the_last_run_started(
    prepared, client_for_prepared
) -> None:
    # No run yet: every shortlisted role is new.
    assert client_for_prepared.get("/api/today").json()["counts"]["new_high_fit"] == 1

    now = datetime.now(UTC)
    with prepared.ctx.repos() as repos:
        repos.events.append(
            "opportunity", prepared.opportunity_id, EventType.JOB_SHORTLISTED, {}, actor="system"
        )
        run = repos.pipeline_runs.create(
            kind="full", prepare=False, candidate_id=prepared.candidate_id, search_profile_id=None
        )
        run.status = "COMPLETED"
        run.started_at = now + timedelta(hours=1)  # began after the role was shortlisted
        run.finished_at = now + timedelta(hours=1)
    body = client_for_prepared.get("/api/today").json()
    assert body["counts"]["new_high_fit"] == 0
    assert body["last_run"]["status"] == "COMPLETED"

    with prepared.ctx.repos() as repos:
        run = repos.pipeline_runs.list_recent(1)[0]
        run.started_at = now - timedelta(days=1)  # began before the role was shortlisted
    assert client_for_prepared.get("/api/today").json()["counts"]["new_high_fit"] == 1

    # A role that has since been passed is not "new" any more.
    pass_opportunity(prepared.ctx, prepared.opportunity_id)
    assert client_for_prepared.get("/api/today").json()["counts"]["new_high_fit"] == 0


def test_employer_responses_count_since_the_last_time_they_were_marked_seen(
    prepared, client_for_prepared
) -> None:
    with prepared.ctx.repos() as repos:
        repos.outcomes.add(
            application_id=prepared.application_id,
            opportunity_id=prepared.opportunity_id,
            kind="PHONE_SCREEN",
        )
        repos.outcomes.add(
            application_id=prepared.application_id,
            opportunity_id=prepared.opportunity_id,
            kind="NO_RESPONSE",
        )
    assert client_for_prepared.get("/api/today").json()["counts"]["employer_responses"] == 1

    seen = client_for_prepared.post("/api/today/seen")
    assert seen.status_code == 200
    body = client_for_prepared.get("/api/today").json()
    assert body["counts"]["employer_responses"] == 0
    assert body["last_seen_at"] == seen.json()["last_seen_at"]

    with prepared.ctx.repos() as repos:
        repos.outcomes.add(
            application_id=prepared.application_id,
            opportunity_id=prepared.opportunity_id,
            kind="INTERVIEW",
        )
    assert client_for_prepared.get("/api/today").json()["counts"]["employer_responses"] == 1


def test_today_shows_the_active_run_and_the_last_finished_one(tmp_path, monkeypatch) -> None:
    use_registry(tmp_path, monkeypatch)
    seed_candidate()
    ctx = EngineContext.create()
    try:
        service = RunService(ctx, discovery_http=mock_http(), research_http=mock_http())
        finished = service.run_sync("rank", False)
        queued = RunService(ctx).create("full")
    finally:
        ctx.close()
    body = TestClient(app).get("/api/today").json()
    assert body["last_run"]["id"] == finished["id"]
    assert body["active_run"]["id"] == queued
    assert body["search_profile_name"] == "Customer success - remote"


def test_the_pages_share_one_shell_with_the_same_navigation() -> None:
    client = TestClient(app)
    paths = [
        "/",
        "/opportunities",
        "/applications",
        "/pipeline",
        "/candidate",
        "/searches",
        "/answers",
        "/settings",
        "/packages/pkg_x",
    ]
    for path in paths:
        html = client.get(path).text
        assert 'class="topnav"' in html, path
        for href in ("/", "/opportunities", "/applications", "/pipeline", "/candidate"):
            assert f'href="{href}"' in html, (path, href)
        assert 'href="/searches"' in html and 'href="/settings"' in html, path
    today = client.get("/").text
    assert 'aria-current="page"' in today
    assert "A clearer path forward." in today
    assert 'id="run-now"' in today and "Run search" in today
    assert "Record what an employer does" in client.get("/pipeline").text
    # The resume panel moved to the Candidate page.
    assert 'id="resume-form"' not in today
    assert 'id="resume-form"' in client.get("/candidate").text


# ---------------------------------------------------------------------------
# Candidate: resume files
# ---------------------------------------------------------------------------


def _docx_bytes(build) -> bytes:
    document = Document()
    build(document)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_docx_text_is_paragraphs_and_table_cells_in_order() -> None:
    def build(document) -> None:
        document.add_paragraph("JORDAN RIVERA")
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Skills"
        table.cell(0, 1).text = "Excel"
        table.cell(1, 0).text = "Phone"
        merged = table.cell(1, 0).merge(table.cell(1, 1))
        merged.text = "555-0100"
        document.add_paragraph("EXPERIENCE")

    text, source_format = extract_resume_text("resume.DOCX", _docx_bytes(build))
    assert source_format == "docx"
    assert text.split("\n") == ["JORDAN RIVERA", "Skills", "Excel", "555-0100", "EXPERIENCE"]


def test_txt_is_utf8_with_a_windows_1252_fallback_and_422_when_neither() -> None:
    assert extract_resume_text("a.txt", "Café manager".encode()) == ("Café manager", "text")
    assert extract_resume_text("a.txt", b"\xef\xbb\xbfplain")[0] == "plain"
    assert extract_resume_text("a.txt", "Café manager".encode("cp1252"))[0] == "Café manager"
    with pytest.raises(ValidationFailure):
        extract_resume_text("a.txt", b"\x81\x8d\x8f")
    with pytest.raises(ValidationFailure):
        extract_resume_text("old.doc", b"whatever")
    with pytest.raises(ValidationFailure):
        extract_resume_text("bad.pdf", b"not a pdf")
    with pytest.raises(ValidationFailure):
        extract_resume_text("bad.docx", b"not a docx")


def test_resume_upload_accepts_docx_and_rejects_the_undecodable(fixture_resume_text) -> None:
    client = TestClient(app)

    def build(document) -> None:
        for line in fixture_resume_text.splitlines():
            document.add_paragraph(line)

    response = client.post(
        "/api/candidate/resume",
        files={
            "resume": (
                "resume.docx",
                _docx_bytes(build),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert response.status_code == 200, response.text
    assert response.json()["source_format"] == "docx"
    assert response.json()["truth_level"] == "INFERRED"

    bad = client.post(
        "/api/candidate/resume", files={"resume": ("resume.txt", b"\x81\x8d\x8f", "text/plain")}
    )
    assert bad.status_code == 422
    empty = client.post(
        "/api/candidate/resume", files={"resume": ("resume.txt", b"   \n", "text/plain")}
    )
    assert empty.status_code == 422

    listing = client.get("/api/candidate").json()
    assert [r["source_format"] for r in listing["resumes"]] == ["docx"]
    assert listing["resumes"][0]["is_master"] is True


def test_resume_upload_stores_a_text_resume_as_inferred(fixture_resume_text) -> None:
    client = TestClient(app)
    response = client.post(
        "/api/candidate/resume",
        files={"resume": ("resume.txt", fixture_resume_text.encode("cp1252"), "text/plain")},
    )
    assert response.status_code == 200, response.text
    sections = {s["title"]: s for s in client.get("/api/candidate").json()["sections"]}
    employment = sections["Employment"]["rows"]
    assert employment and {row["level"] for row in employment} == {"INFERRED"}


# ---------------------------------------------------------------------------
# Candidate: identity
# ---------------------------------------------------------------------------

IDENTITY = {
    "first_name": "Maria",
    "last_name": "Lopez",
    "email": "maria.lopez@example.test",
    "phone": "(214) 555-0199",
    "street_address": "12 Main St",
    "city": "Plano",
    "state": "TX",
    "postal_code": "75024",
    "country": "USA",
    "linkedin_url": "https://www.linkedin.com/in/maria-lopez-example",
    "headline": "Administrative professional",
}


def test_saving_the_identity_form_stores_verified_values_in_a_new_version() -> None:
    client = TestClient(app)
    before = client.get("/api/candidate").json()
    assert before["candidate_id"] is None

    saved = client.put("/api/candidate/identity", json=IDENTITY)
    assert saved.status_code == 200, saved.text
    assert saved.json()["truth_level"] == "VERIFIED"
    body = client.get("/api/candidate").json()
    assert body["identity"] == IDENTITY
    profile = body["profile"]
    assert profile["identity"]["full_name"] == "Maria Lopez"
    for path in ("identity.first_name", "identity.email", "identity.postal_code", "headline"):
        assert body["provenance"][path]["level"] == "VERIFIED", path
    assert profile["answer_memory"]["street_address"]["level"] == "VERIFIED"
    identity_rows = {r["label"]: r for r in body["sections"][0]["rows"]}
    assert identity_rows["Email"]["level"] == "VERIFIED"
    assert identity_rows["Street address"]["level"] == "VERIFIED"

    again = client.put("/api/candidate/identity", json={"phone": "(214) 555-0000"})
    assert again.json()["version"] == saved.json()["version"] + 1
    after = client.get("/api/candidate").json()["identity"]
    assert after["phone"] == "(214) 555-0000" and after["first_name"] == "Maria"

    assert client.put("/api/candidate/identity", json={"email": "not an email"}).status_code == 422
    assert client.put("/api/candidate/identity", json={"shoe_size": "9"}).status_code == 422


def test_owner_identity_wins_over_the_resume_in_either_order(fixture_resume_text) -> None:
    client = TestClient(app)
    client.put("/api/candidate/identity", json=IDENTITY)
    upload = client.post(
        "/api/candidate/resume",
        files={"resume": ("resume.txt", fixture_resume_text.encode(), "text/plain")},
    )
    assert upload.status_code == 200
    body = client.get("/api/candidate").json()
    assert body["profile"]["identity"]["email"] == "maria.lopez@example.test"
    assert body["profile"]["identity"]["full_name"] == "Maria Lopez"
    assert body["provenance"]["identity.email"]["level"] == "VERIFIED"
    assert body["profile"]["headline"] == "Administrative professional"
    # What the resume did add is still inferred.
    assert body["provenance"]["employment"]["level"] == "INFERRED"

    # And the other way round: a resume first, then the owner's statement replaces it.
    other = TestClient(app)
    other.put("/api/candidate/identity", json={"email": "later@example.test"})
    assert other.get("/api/candidate").json()["identity"]["email"] == "later@example.test"


def test_packages_read_the_owner_identity(prepared) -> None:
    save_owner_identity(prepared.ctx, prepared.candidate_id, IDENTITY)
    result = DocumentService(prepared.ctx).prepare_package(prepared.opportunity_id)
    with prepared.ctx.read_repos() as repos:
        payload = repos.packages.get(result.package_id).payload
    assert payload["applicant_first_name"] == "Maria"
    assert payload["applicant_last_name"] == "Lopez"
    assert payload["applicant_email"] == "maria.lopez@example.test"
    assert payload["applicant_postal_code"] == "75024"
    assert payload["applicant_street_address"] == "12 Main St"


# ---------------------------------------------------------------------------
# Pass
# ---------------------------------------------------------------------------


def _opportunity_state(prepared) -> dict:
    with prepared.ctx.read_repos() as repos:
        opportunity = repos.opportunities.get(prepared.opportunity_id)
        job = repos.jobs.get(prepared.job_id)
        application = repos.applications.for_opportunity(prepared.opportunity_id)
        package = repos.packages.latest_for_opportunity(prepared.opportunity_id)
        return {
            "opportunity": opportunity.status,
            "job": job.status,
            "shortlisted": opportunity.shortlisted,
            "reasons": (opportunity.rejection or {}).get("reasons"),
            "application": application.status if application else None,
            "package": package.status if package else None,
        }


def test_pass_rejects_the_role_withdraws_its_unapproved_package_and_leaves_an_event(
    prepared, client_for_prepared
) -> None:
    response = client_for_prepared.post(f"/api/opportunities/{prepared.opportunity_id}/pass")
    assert response.status_code == 200, response.text
    state = _opportunity_state(prepared)
    assert state == {
        "opportunity": "REJECTED",
        "job": "REJECTED",
        "shortlisted": False,
        "reasons": ["OWNER_PASSED"],
        "application": "WITHDRAWN",
        "package": "SUPERSEDED",
    }
    with prepared.ctx.read_repos() as repos:
        events = repos.events.for_entity("opportunity", prepared.opportunity_id)
        rejected = [e for e in events if e.event_type == "JOB_REJECTED"]
        assert rejected and rejected[-1].payload["reasons"] == ["OWNER_PASSED"]
        assert rejected[-1].actor == "owner"
    # Passing twice is harmless.
    again = client_for_prepared.post(f"/api/opportunities/{prepared.opportunity_id}/pass")
    assert again.status_code == 200 and again.json()["already_passed"] is True
    # The pass is no longer on the shortlist the page shows.
    assert client_for_prepared.get("/api/opportunities").json()["opportunities"] == []


def test_pass_is_refused_once_the_application_is_past_awaiting_approval(
    prepared, client_for_prepared
) -> None:
    from sabra_engine.approval import ApprovalService

    ApprovalService(prepared.ctx).approve(
        prepared.package_id, approved_by="test", acknowledge_unresolved=True
    )
    response = client_for_prepared.post(f"/api/opportunities/{prepared.opportunity_id}/pass")
    assert response.status_code == 409, response.text
    assert "past awaiting approval" in response.json()["detail"]
    state = _opportunity_state(prepared)
    assert state["opportunity"] == "APPROVED"
    assert state["application"] == "APPROVED"
    assert state["reasons"] is None
    assert client_for_prepared.post("/api/opportunities/opp_nope/pass").status_code == 404


def test_a_ranking_pass_never_brings_a_passed_role_back(prepared) -> None:
    pass_opportunity(prepared.ctx, prepared.opportunity_id)
    for rescore in (False, True):
        result = IntelligenceService(prepared.ctx).run(
            candidate_id=prepared.candidate_id,
            search_profile_id=prepared.search_profile_id,
            rescore=rescore,
        )
        assert result.shortlisted == 0
        assert _opportunity_state(prepared)["opportunity"] == "REJECTED"
        assert _opportunity_state(prepared)["shortlisted"] is False


def test_pass_works_on_a_role_with_no_package_yet(tmp_path, monkeypatch) -> None:
    use_registry(tmp_path, monkeypatch)
    seed_candidate()
    ctx = EngineContext.create()
    try:
        record = RunService(ctx, discovery_http=mock_http(), research_http=mock_http()).run_sync(
            "full", False
        )
        with ctx.read_repos() as repos:
            opportunity = repos.opportunities.shortlist(record["candidate_id"])[0]
            assert opportunity.status == "RESEARCHED"
        passed = pass_opportunity(ctx, opportunity.id)
        assert passed["already_passed"] is False
        with ctx.read_repos() as repos:
            assert repos.opportunities.get(opportunity.id).status == "REJECTED"
            assert repos.applications.for_opportunity(opportunity.id) is None
    finally:
        ctx.close()


# ---------------------------------------------------------------------------
# the opportunity card
# ---------------------------------------------------------------------------


def test_the_card_carries_pay_confidence_and_the_primary_gap(
    prepared, client_for_prepared
) -> None:
    with prepared.ctx.repos() as repos:
        repos.fit.upsert(
            prepared.opportunity_id,
            prepared.candidate_id,
            {
                "fit_score": 0.8,
                "confidence": 0.72,
                "gaps": ["No CRM experience shown", "Territory is large"],
                "why_this_job": "Strong overlap",
            },
        )
    row = client_for_prepared.get("/api/opportunities").json()["opportunities"][0]
    assert row["compensation_text"] == "$110,000 to $140,000 a year"
    assert row["confidence"] == pytest.approx(0.72)
    assert row["primary_gap"] == "No CRM experience shown"
    assert row["gaps"][1] == "Territory is large"

    with prepared.ctx.repos() as repos:
        job = repos.jobs.get(prepared.job_id)
        job.salary_min = job.salary_max = None
    row = client_for_prepared.get("/api/opportunities").json()["opportunities"][0]
    assert row["compensation_text"] == "not listed"


def test_the_list_follows_the_active_search_unless_asked_for_all(
    prepared, client_for_prepared
) -> None:
    other = {"name": "Another search"}
    created = client_for_prepared.post("/api/searches", json=other).json()["profile"]["id"]
    assert len(client_for_prepared.get("/api/opportunities").json()["opportunities"]) == 1
    client_for_prepared.post(f"/api/searches/{created}/activate")
    assert client_for_prepared.get("/api/opportunities").json()["opportunities"] == []
    everything = client_for_prepared.get("/api/opportunities?all_profiles=true").json()
    assert len(everything["opportunities"]) == 1


# ---------------------------------------------------------------------------
# Settings
# ---------------------------------------------------------------------------


def test_settings_reports_the_running_configuration(monkeypatch) -> None:
    from sabra_domain.settings import reset_settings_cache

    monkeypatch.setenv("SABRA_LLM_PROVIDER", "anthropic")
    monkeypatch.setenv("SABRA_LLM_MODEL", "some-model")
    monkeypatch.setenv("SABRA_HANDOFF_TIMEOUT_SECONDS", "600")
    reset_settings_cache()
    seed_candidate()
    body = TestClient(app).get("/api/settings").json()
    assert body["automation_dry_run"] is True
    assert body["dry_run_env_var"] == "SABRA_AUTOMATION_DRY_RUN"
    assert body["llm_provider"] == "anthropic" and body["llm_model"] == "some-model"
    assert body["handoff_timeout_seconds"] == 600
    assert "isolated.db" in body["database_url"]
    assert body["artifacts_dir"].endswith("artifacts")
    assert body["active_search_profile_name"] == "Customer success - remote"
    assert isinstance(body["playwright"]["chromium_installed"], bool)
    assert isinstance(body["playwright"]["package_installed"], bool)
    import tomllib
    from pathlib import Path

    pyproject = tomllib.loads((Path(__file__).parents[1] / "pyproject.toml").read_text())
    assert body["version"] == pyproject["project"]["version"]
    assert "Settings" in TestClient(app).get("/settings").text


def test_playwright_status_looks_for_chromium_without_launching(tmp_path, monkeypatch) -> None:
    from sabra_api.main import _playwright_status

    monkeypatch.setenv("PLAYWRIGHT_BROWSERS_PATH", str(tmp_path))
    assert _playwright_status()["chromium_installed"] is False
    (tmp_path / "chromium-1234").mkdir()
    status = _playwright_status()
    assert status["chromium_installed"] is True
    assert status["browsers_path"] == str(tmp_path)
